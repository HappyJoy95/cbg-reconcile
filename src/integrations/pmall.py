# -*- coding: utf-8 -*-
"""pmall（华为买家中心）接口客户端 —— 「价格及返利政策」导出手动更新链路。

## 四个端点（2026-09-29 全部实测通过，见开发目标文档）

    换 csrf   POST /phoenix-gws/phoenix.sso.csrf.token
              ⚠ 身份必须是 uid；换出来的 `GUEST` 打不动业务接口（实测 401/空）
    提交导出  POST mp.asyncTask.export
              ⚠ `taskTypeDisplay` 必须是「价格返利政策商品导出」——
                curl 里那 10 个空格会回 `params error`（踩过）
    任务列表  GET  mp.task.myExportList?siteId&pageIndex&pageSize…
              响应是 `{pageArgs, dataList}`，下载凭 `dataList[].id`（不是 taskId）
    下载      GET  mp.task.downloadFromEDM?taskId=<id>
              回 xlsx 字节流；`PK` 魔数校验，错误页不许冒充

页面入口：`#/group/pmall-cn/my-imExport?tab=myExport`（我的导入导出 › 导出查询）。

## 会话策略（用户 2026-09-29 拍板的「A 案」）

    ① 复用**还开着的受控窗口** —— 我们自己 launch 的 profile + 调试端口，
       登录成功后**不关窗口**（关了 session cookie 会被 Chrome 清掉，实测）；
       端口记在 `.secrets/pmall-window.json`。
    ② 退 cookie jar（`.secrets/pmall-session.json`）—— 服务端会话空闲约
       半小时会掉（实测 12:24 存、13:08 已 GUEST），所以拿到就现换 csrf 验身份。
    ③ 弹可见窗口**人工登录** —— uniportal 有短信二次验证
       （`login-doubleFactor.html`，自动填表过不去），**人工操作一次**，
       登完窗口留着，下次走 ①。

⚠ **隐私边界**（`browser.py` 顶部 / AGENTS 坑 9 —— 设计约束，不是保守）：
   只碰**我们自己启动的** profile，**不读用户日常浏览器的 cookie**。
   （用户问过"用我正常浏览器是不是就免验证码"：技术上接不上——没开调试端口
   + Chromium 单实例；免验证码的上限两边完全一样，都只取决于会话是否活着。）

⚠ **不做定时抓取**（同日拍板）：手动触发；本模块只管"拿到数据"。
⚠ 撞到验证码/登录报错**立刻停手不重试** —— 连打失败会锁号（同 `try_auto_login` 的规矩）。
"""

from __future__ import annotations

import base64
import io
import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple
from urllib.parse import urlencode

import requests

from . import browser as _browser
from .cdp import CdpError, http_json
from ..paths import ROOT

# --------------------------------------------------------------------- 常量
BASE = "https://pmall.huawei.com"
#: 低代码平台的站点 id = 请求头里的 `pmall-app-id`（全站通用）
SITE_ID = "af5b389dc0d64249bbcffb83962653f0"
#: ⚠ **本店（胶州机场店）的店铺 id** —— 换店要传自己的（页面从 bulkQueryShop 拿）
DEFAULT_SHOP_ID = "ac7e1d6a10434c3e9f549b58b988621b"

TASK_TYPE = "pmall-china-PriceRebatePolicyExport"
#: ⚠ 真文案，不是空格 —— 这是 `params error` 的坑源
TASK_DISPLAY = "价格返利政策商品导出"
SERVICE_ID = "ideal.mygoods.export"

EXPORT_PATH = "/api-gateway/services/mp.asyncTask.export"
LIST_PATH = "/api-gateway/services/mp.task.myExportList"
DOWNLOAD_PATH = "/api-gateway/services/mp.task.downloadFromEDM"
CSRF_PATH = "/phoenix-gws/phoenix.sso.csrf.token"

#: 从价格页登录会自己跳这里；登录完跳回价格页
PMALL_URL = BASE + "/#/group/pmall-cn/price-and-rebate-policy"
IMEXPORT_PAGE = "my-imExport"

STATUS_OK = "success"
STATUS_FAIL = frozenset(("failed", "fail", "error"))

WINDOW_FILE = "pmall-window.json"
JAR_FILE = "pmall-session.json"

USER_AGENT = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/154.0.0.0 Safari/537.36 Edg/154.0.0.0")


class PmallError(RuntimeError):
    """pmall 链路的问题（会话 / 提交 / 轮询 / 下载 / 解析）。

    ⚠ **别用 SystemExit** —— 这个模块会被进程内调用（AGENTS 坑 11）。
    """


# ------------------------------------------------------------------ 纯函数
def identity(csrf: str) -> Dict[str, object]:
    """解 csrf JWT 的 payload（**纯函数**，坏了回空 dict 不抛）。"""
    try:
        parts = str(csrf or "").split(".")
        pad = parts[1] + "=" * (-len(parts[1]) % 4)
        got = json.loads(base64.urlsafe_b64decode(pad))
        return got if isinstance(got, dict) else {}
    except (IndexError, ValueError, TypeError):
        return {}


def uid_of(csrf: str) -> str:
    """csrf 身份 —— **`GUEST` 当没登录**（换回来的匿名 token 就是它，实测打不动接口）。"""
    who = str(identity(csrf).get("CSRF") or "").strip()
    return "" if who in ("", "GUEST") else who


# ------------------------------------------------------------- 会话值对象
@dataclass
class PmallSession:
    cookies: str
    csrf: str
    source: str = ""            # window / jar / login

    @property
    def uid(self) -> str:
        return uid_of(self.csrf)

    def headers(self, page: str = IMEXPORT_PAGE) -> Dict[str, str]:
        """业务接口的公共请求头（四个端点实测够用；`sw8` 那些追踪头不需要）。"""
        return {
            "accept": "application/json",
            "cookie": self.cookies,
            "user-agent": USER_AGENT,
            "referer": BASE + "/",
            "origin": BASE,
            "content-type": "application/json;charset=UTF-8",
            "x-pix-csrf-token": self.csrf or "",
            "x-csrf-token": "null",
            "x-pix-app-type": "APP",
            "x-pix-app-url": "pmall-cn",
            "x-pix-product-id": "phoenix",
            "x-pix-tenant-id": "phoenix",
            "x-pix-page-url": page,
            "pmall-app-id": SITE_ID,
            "pmall-user-lang": "zh_CN",
        }


# --------------------------------------------------------- 薄 HTTP 壳（打桩点）
def _post(url: str, *, headers: Dict[str, str], data=None, timeout: float = 30):
    return requests.post(url, headers=headers, data=data, timeout=timeout)


def _get(url: str, *, headers: Dict[str, str], timeout: float = 30):
    return requests.get(url, headers=headers, timeout=timeout)


#: 轮询/等待的 sleep —— 单独一个名字，测试打掉它就不用真等
_sleep = time.sleep


def _json(resp) -> dict:
    try:
        body = resp.json()
    except ValueError:
        return {}
    return body if isinstance(body, dict) else {}


def _snippet(obj) -> str:
    try:
        return json.dumps(obj, ensure_ascii=False)[:200]
    except (TypeError, ValueError):
        return str(obj)[:200]


def _noop(_msg: str) -> None:
    pass


# --------------------------------------------------------------- 状态文件
def _secrets_dir(root=None) -> Path:
    return Path(root or ROOT) / ".secrets"


def _read_json(path: Path) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def _write_json(path: Path, data: dict) -> bool:
    """原子写（tmp + replace）；失败回 False，**不抛**（状态文件不该打断主流程）。"""
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(path.suffix + ".tmp")
        tmp.write_text(json.dumps(data or {}, ensure_ascii=False, indent=2),
                       encoding="utf-8")
        tmp.replace(path)
        return True
    except OSError:
        return False


def load_window(root=None) -> dict:
    """受控登录窗口的登记（`{port, profile, pid, opened_at}`）；没有就 `{}`。"""
    return _read_json(_secrets_dir(root) / WINDOW_FILE)


def save_window(root, info: dict) -> bool:
    return _write_json(_secrets_dir(root) / WINDOW_FILE, info or {})


def load_jar(root=None) -> dict:
    """cookie jar（`{cookies, csrf, source, saved_at}`）；没有就 `{}`。"""
    return _read_json(_secrets_dir(root) / JAR_FILE)


def save_jar(root, sess: "PmallSession") -> bool:
    return _write_json(_secrets_dir(root) / JAR_FILE, {
        "cookies": sess.cookies, "csrf": sess.csrf, "source": sess.source,
        "saved_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    })


# ---------------------------------------------------------------- 浏览器侧
def _alive(port: int) -> bool:
    """调试端口还活着吗（浏览器被关 = 端口没了）。"""
    try:
        http_json(int(port), "/json/version", timeout=1)
        return True
    except (CdpError, OSError, ValueError):
        return False


def exchange_csrf(cookies: str, *, timeout: float = 20) -> Optional[str]:
    """cookie → 新 csrf（每次现换，别落盘复用 —— token 才半小时命）。"""
    if not cookies:
        return None
    headers = {
        "accept": "application/json",
        "cookie": cookies,
        "user-agent": USER_AGENT,
        "referer": BASE + "/",
        "origin": BASE,
        "content-type": "application/json;charset=UTF-8",
    }
    try:
        resp = _post(BASE + CSRF_PATH, headers=headers, data="{}", timeout=timeout)
        if resp.status_code != 200:
            return None
        data = _json(resp).get("data")
        return data if isinstance(data, str) and data else None
    except (requests.RequestException, ValueError):
        return None


def probe_port(port: int, *, source: str = "window") -> Optional[PmallSession]:
    """从受控窗口收一套 (cookie + csrf)。收不齐回 None（**不抛**）。"""
    try:
        cookies, _detail = _browser.cookies_from_browser(int(port))
    except (CdpError, OSError, ValueError):
        return None
    if not cookies:
        return None
    csrf = exchange_csrf(cookies)
    if not csrf:
        return None
    return PmallSession(cookies=cookies, csrf=csrf, source=source)


def session_from_jar(root=None) -> Optional[PmallSession]:
    """jar 里的 cookie 现换一次 csrf —— **身份是 uid 才算数**（过期的 jar 会换出 GUEST）。"""
    jar = load_jar(root)
    cookies = str(jar.get("cookies") or "")
    if not cookies:
        return None
    csrf = exchange_csrf(cookies)
    sess = PmallSession(cookies=cookies, csrf=csrf or "", source="jar")
    return sess if sess.uid else None


# --------------------------------------------------------------- 登录三分支
def wait_login_in(port: int, *, root=None, timeout: float = 600.0,
                  say: Optional[Callable[[str], None]] = None) -> PmallSession:
    """窗口还开着但会话过期 → **同一个窗口**里请人重新登录（可能要短信）。"""
    say = say or _noop
    say("窗口还开着但会话过期了 —— 请在那个窗口重新登录（可能要短信验证）")
    _browser.goto_url(int(port), PMALL_URL)
    end = time.time() + timeout
    while time.time() < end:
        if not _alive(port):
            raise PmallError("登录窗口被关掉了 —— 会话没拿到")
        sess = probe_port(port, source="window")
        if sess and sess.uid:
            say("登录完成（%s）—— 窗口先留着，下次直接复用" % sess.uid)
            save_jar(root, sess)
            return sess
        _sleep(3.0)
    raise PmallError("等登录超时（%d 秒）没完成 —— 下次点更新会重新弹窗" % int(timeout))


def launch_login(root=None, *, timeout: float = 600.0,
                 say: Optional[Callable[[str], None]] = None,
                 headless: bool = False) -> PmallSession:
    """弹可见窗口人工登录（密码 + 短信）。**成功后不关窗口** —— 那就是 ① 的活窗。"""
    say = say or _noop
    profile = _secrets_dir(root) / "pmall-profile"
    profile.mkdir(parents=True, exist_ok=True)
    proc, port = _browser.launch(profile, url=PMALL_URL, headless=headless)
    save_window(root, {"port": int(port), "profile": str(profile),
                       "pid": getattr(proc, "pid", 0),
                       "opened_at": time.strftime("%Y-%m-%d %H:%M:%S")})
    mins = max(1, int(timeout) // 60)
    say("已打开登录窗口 —— 密码/短信请人工完成（最多 %d 分钟；登完别关窗口）" % mins)
    end = time.time() + timeout
    while time.time() < end:
        if proc.poll() is not None or not _alive(port):
            raise PmallError("登录窗口被关掉了（退出码 %s）—— 会话没拿到"
                             % getattr(proc, "returncode", "?"))
        sess = probe_port(port, source="window")
        if sess and sess.uid:
            say("登录完成（%s）—— 窗口先留着，下次直接复用" % sess.uid)
            save_jar(root, sess)
            return sess
        _sleep(3.0)
    raise PmallError("等登录超时（%d 秒）没完成 —— 下次点更新会重新弹窗" % int(timeout))


def ensure_session(root=None, *, timeout: float = 600.0,
                   say: Optional[Callable[[str], None]] = None) -> PmallSession:
    """**A 案三分支**：活窗 → jar → 弹窗人工。拿不到就 PmallError（文案说人话）。"""
    say = say or _noop
    try:
        port = int(load_window(root).get("port") or 0)
    except (TypeError, ValueError):
        port = 0
    if port and _alive(port):
        sess = probe_port(port, source="window")
        if sess and sess.uid:
            say("复用已开着的登录窗口（%s）" % sess.uid)
            save_jar(root, sess)
            return sess
        return wait_login_in(port, root=root, timeout=timeout, say=say)
    sess = session_from_jar(root)
    if sess:
        say("复用已存会话（%s）" % sess.uid)
        save_jar(root, sess)
        return sess
    say("没有可用会话 —— 打开登录窗口（密码/短信请人工操作）")
    return launch_login(root, timeout=timeout, say=say)


# ------------------------------------------------------------------ 四端点
def export_policy(sess: PmallSession, *, shop_id: str = DEFAULT_SHOP_ID,
                  uid: Optional[str] = None, timeout: float = 30) -> dict:
    """提交「价格返利政策商品导出」。uid 取自 csrf 身份（不信传参也行）。"""
    who = (uid or "").strip() or sess.uid
    if not who:
        raise PmallError("会话不是登录态（csrf 身份是 GUEST）—— 先走 ensure_session()")
    payload = {
        "siteId": SITE_ID,
        "userId": who,
        "taskType": TASK_TYPE,
        "taskTypeDisplay": TASK_DISPLAY,          # ⚠ 真文案，空格会 params error
        "serviceId": SERVICE_ID,
        "extendParams": {"uid": who, "lang": "zh_CN", "shopId": shop_id,
                         "skuName": "", "goodsCode": "", "spuName": ""},
    }
    resp = _post(BASE + EXPORT_PATH, headers=sess.headers(IMEXPORT_PAGE),
                 data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                 timeout=timeout)
    if resp.status_code != 200:
        raise PmallError("导出提交 HTTP %s：%s" % (resp.status_code, _snippet(_json(resp))))
    body = _json(resp)
    if str(body.get("code")) != "200":
        raise PmallError("导出提交被拒：%s" % _snippet(body))
    return body


def list_exports(sess: PmallSession, *, page_index: int = 1, page_size: int = 10,
                 timeout: float = 30) -> List[dict]:
    """任务列表（我的导入导出 › 导出查询）。响应 `{pageArgs, dataList}`。"""
    query = urlencode({"siteId": SITE_ID, "pageIndex": page_index,
                       "pageSize": page_size, "fileName": "", "taskType": "",
                       "status": "", "createdDateStart": "", "createdDatend": "",
                       "t": int(time.time() * 1000)})
    resp = _get(BASE + LIST_PATH + "?" + query,
                headers=sess.headers(IMEXPORT_PAGE), timeout=timeout)
    if resp.status_code != 200:
        raise PmallError("任务列表 HTTP %s" % resp.status_code)
    body = _json(resp)
    rows = body.get("dataList")
    if not isinstance(rows, list):
        raise PmallError("任务列表读不出来：%s" % _snippet(body))
    return rows


def wait_new_export(sess: PmallSession, *, known, timeout: float = 120.0,
                    interval: float = 3.0, page_size: int = 10,
                    say: Optional[Callable[[str], None]] = None) -> dict:
    """等**新出现**的任务跑成功（`known` = 提交前记下的 id 集合）。

    失败当场报错；超时报错并指路页面 —— 不空转到天荒地老。
    """
    say = say or _noop
    known = set(known or ())
    end = time.time() + timeout
    while True:
        fresh = [r for r in list_exports(sess, page_size=page_size)
                 if r.get("taskType") == TASK_TYPE and r.get("id")
                 and r.get("id") not in known]
        if fresh:
            fresh.sort(key=lambda r: str(r.get("createdDate") or ""), reverse=True)
            rec = fresh[0]
            status = str(rec.get("status") or "")
            if status == STATUS_OK:
                say("导出完成：%s" % rec.get("fileName"))
                return rec
            if status in STATUS_FAIL:
                raise PmallError("导出任务失败（%s）：%s"
                                 % (status, rec.get("fileName") or rec.get("id")))
        if time.time() >= end:
            raise PmallError("等导出结果超时（%d 秒）——可到「我的导入导出 › 导出查询」"
                             "手动看一眼" % int(timeout))
        say("导出处理中…")
        _sleep(interval)


def download(sess: PmallSession, task_id: str, *, timeout: float = 60) -> bytes:
    """下载任务产物；**`PK` 魔数硬校验** —— 登录态失效时回来的是 HTML 错误页。"""
    task_id = str(task_id or "").strip()
    if not task_id:
        raise PmallError("下载缺 taskId")
    url = BASE + DOWNLOAD_PATH + "?" + urlencode({"taskId": task_id})
    headers = sess.headers(IMEXPORT_PAGE)
    headers["accept"] = "*/*"
    resp = _get(url, headers=headers, timeout=timeout)
    if resp.status_code != 200:
        raise PmallError("下载 HTTP %s" % resp.status_code)
    data = resp.content or b""
    if data[:2] != b"PK":
        raise PmallError("下载到的不是 xlsx（%d 字节，开头 %r）——多半是登录态失效的错误页"
                         % (len(data), data[:16]))
    return data


def fetch_policy(sess: PmallSession, *, shop_id: str = DEFAULT_SHOP_ID,
                 timeout: float = 120.0,
                 say: Optional[Callable[[str], None]] = None) -> Tuple[bytes, dict]:
    """一条龙：记旧任务 → 提交 → 等新任务成功 → 下载。返回 (xlsx字节, 任务记录)。"""
    say = say or _noop
    known = {r.get("id") for r in list_exports(sess, page_size=10)}
    export_policy(sess, shop_id=shop_id)
    say("已提交导出任务，等它跑完…")
    rec = wait_new_export(sess, known=known, timeout=timeout, say=say)
    data = download(sess, rec["id"])
    say("下载完成：%d 字节" % len(data))
    return data, rec


# ------------------------------------------------------------------ 解析
#: 政策表的商品键（xlsx 表头原样，9 列见开发目标文档）
POLICY_CODE_FIELD = "商品编码"


def parse_policy(data: bytes) -> List[Dict[str, object]]:
    """xlsx → 行 dict 列表。**表头原样保留**（含 `基准提货价*` 的星号）。"""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:                                     # noqa: BLE001
        raise PmallError("xlsx 打不开：%s: %s" % (type(e).__name__, e))
    try:
        sheets = wb.sheetnames
        if not sheets:
            raise PmallError("xlsx 没有 sheet")
        ws = wb[sheets[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not any(header):
            raise PmallError("xlsx 第一行不是表头（空文件？）")
        names = [str(c).strip() if c is not None else "" for c in header]
        out: List[Dict[str, object]] = []
        for row in rows:
            if not row or not any(v not in (None, "") for v in row):
                continue                                     # 空行跳过
            item = {}
            for i, val in enumerate(row):
                if i < len(names) and names[i]:
                    item[names[i]] = val
            out.append(item)
        return out
    finally:
        try:
            wb.close()
        except Exception:                                     # noqa: BLE001
            pass


__all__ = [
    "PmallError", "PmallSession", "identity", "uid_of",
    "load_window", "save_window", "load_jar", "save_jar",
    "exchange_csrf", "probe_port", "session_from_jar",
    "wait_login_in", "launch_login", "ensure_session",
    "export_policy", "list_exports", "wait_new_export", "download",
    "fetch_policy", "parse_policy",
    "DEFAULT_SHOP_ID", "TASK_TYPE", "TASK_DISPLAY", "PMALL_URL",
    "POLICY_CODE_FIELD",
]
